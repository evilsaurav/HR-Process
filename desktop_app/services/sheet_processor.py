import difflib
import os
import re
import datetime
import pandas as pd
from config import REQUIRED_CANDIDATE_COLUMNS
from database import db_manager


RECOGNIZED_HEADER_KEYWORDS = {
    # Keka fields
    "employeenumber", "employeecode", "empid", "firstname", "middlename", "lastname",
    "displayname", "fullname", "workemail", "personalemail", "mobilephone", "mobile",
    "dateofbirth", "gender", "maritalstatus", "marriagedate", "bloodgroup",
    "fathername", "mothername", "spousename", "location", "department", "subdepartment",
    "jobtitle", "datejoined", "pannumber", "aadhaarnumber", "pfnumber", "uannumber",
    "employmentstatus", "exitdate", "exitstatus", "terminationtype", "terminationreason",
    "reportingto", "reportingmanageremployeenumber", "leaveplan", "costcenter",
    "workertype", "timetype", "noticeperiod", "legalentity", "businessunit",
    "currentaddress", "permanentaddress", "currentaddressline1", "permanentaddressline1",
    # Joining Process pipeline fields
    "candidatename", "applieddesignation", "negotiatedsalary", "recentdesignation",
    "highesteducationqualification", "highesteducation", "totalyearsofexperience",
    "totalexperience", "coretechnicalskills", "keyresponsibilities", "contactlocation",
    "email", "status"
}


def read_sheet(file_path: str) -> pd.DataFrame:
    """
    Intelligently reads an Excel (.xlsx, .xls) or CSV file into a pandas DataFrame:
    - Auto-detects header row offset (e.g. Row 2 in Keka reports with company/title banners).
    - Auto-selects best sheet if workbook contains multiple tabs.
    - Auto-detects CSV delimiters (comma, tab, semicolon) and handles encodings.
    """
    if not file_path or not os.path.exists(file_path):
        raise FileNotFoundError(f"File not found: {file_path}")

    def _tok(v):
        return re.sub(r"[^a-z0-9]", "", str(v or "").lower())

    if file_path.lower().endswith((".xlsx", ".xls")):
        with pd.ExcelFile(file_path) as xl:
            best_sheet = xl.sheet_names[0]
            best_row = 0
            max_score = 0

            for s in xl.sheet_names:
                try:
                    df_probe = xl.parse(s, header=None, nrows=15)
                    for r_idx, row in df_probe.iterrows():
                        score = sum(1 for val in row if _tok(val) in RECOGNIZED_HEADER_KEYWORDS)
                        if score > max_score:
                            max_score = score
                            best_row = r_idx
                            best_sheet = s
                except Exception:
                    continue

            header_arg = best_row if max_score >= 2 else 0
            return xl.parse(best_sheet, header=header_arg)
    else:
        # CSV / TSV reader
        for enc in ["utf-8-sig", "utf-8", "cp1252", "latin1"]:
            try:
                with open(file_path, "r", encoding=enc, errors="replace") as f:
                    lines = [f.readline() for _ in range(15)]
                lines = [l for l in lines if l.strip()]
                if not lines:
                    return pd.read_csv(file_path, encoding=enc)

                best_sep = ","
                best_row = 0
                max_score = 0
                for sep in [",", "\t", ";"]:
                    for r_idx, line in enumerate(lines):
                        tokens = [_tok(t.strip()) for t in line.split(sep)]
                        score = sum(1 for t in tokens if t in RECOGNIZED_HEADER_KEYWORDS)
                        if score > max_score:
                            max_score = score
                            best_row = r_idx
                            best_sep = sep

                header_arg = best_row if max_score >= 2 else 0
                return pd.read_csv(file_path, sep=best_sep, header=header_arg, encoding=enc)
            except Exception:
                continue
        return pd.read_csv(file_path)


def validate_columns(df: pd.DataFrame) -> list[str]:
    """Returns a list of required column names missing from the DataFrame."""
    missing = [c for c in REQUIRED_CANDIDATE_COLUMNS if c not in df.columns]
    return missing


def normalize_mobile(value) -> str:
    """Extracts digits and normalizes to the last 10 digits."""
    digits = re.sub(r"\D", "", str(value))
    return digits[-10:] if len(digits) >= 10 else digits


def normalize_name_for_comparison(name: str) -> str:
    """
    Normalizes candidate name for fuzzy similarity matching:
    - Strips honorific prefixes (Mr., Ms., Dr., etc.)
    - Removes punctuation and symbols
    - Converts to lowercase and collapses multiple spaces
    """
    if not name:
        return ""
    cleaned = re.sub(r"^(mr\.|mrs\.|ms\.|dr\.|shri|smt\.)\s+", "", str(name).strip(), flags=re.IGNORECASE)
    cleaned = re.sub(r"[^\w\s]", "", cleaned.lower())
    return " ".join(cleaned.split())


def compute_name_similarity(name1: str, name2: str) -> float:
    """
    Computes candidate name similarity ratio using difflib.SequenceMatcher.
    Returns a float from 0.0 (no similarity) to 1.0 (exact match).
    """
    n1 = normalize_name_for_comparison(name1)
    n2 = normalize_name_for_comparison(name2)
    if not n1 or not n2:
        return 0.0
    if n1 == n2:
        return 1.0
    return round(difflib.SequenceMatcher(None, n1, n2).ratio(), 4)


CANDIDATE_HEADER_SYNONYMS = {
    "candidate_name": ["candidatename", "name", "fullname", "applicantname", "applicant"],
    "mobile": ["mobile", "mobilephone", "phone", "contactnumber", "contactno", "contact", "cell"],
    "email": ["email", "emailid", "mail", "workemail", "personalemail"],
    "applied_designation": ["applieddesignation", "appliedpost", "post", "designation", "jobtitle", "role", "position"],
    "negotiated_salary": ["negotiatedsalary", "salary", "expectedsalary", "currentsalary", "offeredsalary", "grosssalary"],
    "contact_location": ["contactlocation", "appliedlocation", "location", "city", "postinglocation", "address", "state"],
    "recent_designation": ["recentdesignation", "currentdesignation", "experiencedetails", "previousdesignation"],
    "highest_education": ["highesteducationqualification", "highesteducationqualificatio", "highesteducation", "qualification", "education"],
    "total_experience": ["totalyearsofexperienceinferred", "totalyearsofexperienceinf", "totalyearsofexperience", "totalexperience", "experience", "yearsofexperience"],
    "core_technical_skills": ["coretechnicalskills", "technicalskills", "skills", "keyskills"],
    "key_responsibilities": ["keyresponsibilities", "responsibilities", "worksummary", "jobprofile"],
}


def clean_text_value(v) -> str:
    """Cleans empty, null, placeholder strings and strips float .0 suffixes."""
    if v is None:
        return ""
    s = str(v).strip()
    if s.lower() in [
        "none", "nan", "null", "not in source", "not mentioned in source",
        "not specified in resume", "not mentioned in resume", "not mentioned",
        "not specified", "na", "n/a", "--", "-"
    ]:
        return ""
    # Strip .0 on numbers (e.g. 12000.0 -> 12000)
    if s.endswith(".0") and s[:-2].replace(".", "").isdigit():
        return s[:-2]
    return s


def clean_mobile_value(v) -> str:
    """
    Extracts a clean 10-digit mobile number, handling float .0 suffixes,
    multi-number slashes/commas ('9110160907/7209381849'), and +91 prefixes.
    """
    if not v:
        return ""
    s = str(v).strip()
    if s.endswith(".0") and s[:-2].replace(".", "").isdigit():
        s = s[:-2]
    # Check for 10-digit Indian mobile pattern
    m = re.search(r"(?:(?:\+|0{0,2})91[\s-]*)?([6-9]\d{9})", s)
    if m:
        return m.group(1)
    digits = re.sub(r"\D", "", s)
    return digits[-10:] if len(digits) >= 10 else digits


def extract_candidate_rows_from_sheet(file_path: str) -> dict:
    """
    Intelligently reads and extracts candidate rows across ALL sheets in an Excel or CSV file.
    - Dynamically auto-detects header rows and column mappings in multiple stacked batches.
    - Never stops if an optional column is missing; extracts all available fields.
    - Automatically cleans float suffixes, extracts valid 10-digit mobiles, and filters out banner rows.
    - Sets department/project context from sheet names.

    Returns:
      {
          "success": bool,
          "error": str or None,
          "total_sheet_rows": int,
          "valid_rows": list[dict],
          "sheets_processed": list[str],
          "errors": list[str]
      }
    """
    if not file_path or not os.path.exists(file_path):
        return {
            "success": False,
            "error": "The selected file could not be found. Please check the file path and try again.",
            "total_sheet_rows": 0,
            "valid_rows": [],
            "sheets_processed": [],
            "errors": [],
        }

    def _tok(val):
        return re.sub(r"[^a-z0-9]", "", str(val or "").lower())

    BANNER_PHRASES = [
        "interview date", "candidate comparison", "candidate profile",
        "roster", "interview", "plain text values", "summary",
        "candidate professional profile"
    ]

    valid_rows = []
    errors = []
    sheets_processed = []

    # Strategy 1: Excel Multi-Sheet Extraction
    if file_path.lower().endswith((".xlsx", ".xls")):
        try:
            import openpyxl
            wb = openpyxl.load_workbook(file_path, data_only=True)
            sheet_names = wb.sheetnames
        except Exception as ex:
            print("openpyxl load error, attempting pandas fallback:", ex)
            try:
                with pd.ExcelFile(file_path) as xl:
                    sheet_names = xl.sheet_names
            except Exception as e:
                return {
                    "success": False,
                    "error": "We couldn't read this file. Please make sure it's a valid Excel file, close it if open in Microsoft Excel, and try again.",
                    "total_sheet_rows": 0,
                    "valid_rows": [],
                    "sheets_processed": [],
                    "errors": [str(e)],
                }
            wb = None

        for sheet_name in sheet_names:
            try:
                if wb is not None:
                    sheet = wb[sheet_name]
                    raw_rows = list(sheet.iter_rows(values_only=True))
                else:
                    with pd.ExcelFile(file_path) as xl:
                        raw_df = xl.parse(sheet_name, header=None)
                        raw_rows = raw_df.values.tolist()
            except Exception as e:
                errors.append(f"Sheet '{sheet_name}' read error: {e}")
                continue

            # Skip empty sheets
            if not raw_rows or not any(any(r) for r in raw_rows):
                continue

            sheets_processed.append(sheet_name)
            current_col_indices = {}

            for r_idx, row in enumerate(raw_rows):
                if not any(row):
                    continue

                tok_row = [_tok(c) for c in row]
                matches = 0
                temp_map = {}

                # Check if this row is a header row
                for col_idx, t in enumerate(tok_row):
                    if not t:
                        continue
                    for field, keywords in CANDIDATE_HEADER_SYNONYMS.items():
                        if field not in temp_map:
                            if any(t == kw or t.startswith(kw) or kw.startswith(t) for kw in keywords if len(t) >= 4):
                                temp_map[field] = col_idx
                                matches += 1
                                break

                # Header detection threshold: Has candidate_name and (mobile or email or matches >= 3)
                if "candidate_name" in temp_map and ("mobile" in temp_map or "email" in temp_map or matches >= 3):
                    current_col_indices = temp_map
                    continue

                if not current_col_indices or "candidate_name" not in current_col_indices:
                    continue

                name_col = current_col_indices.get("candidate_name")
                raw_name = row[name_col] if name_col is not None and name_col < len(row) else None
                s_name = clean_text_value(raw_name)

                # Skip header repetitions and banner rows
                if not s_name or any(b in s_name.lower() for b in BANNER_PHRASES):
                    continue

                cand_data = {
                    "row_index": r_idx + 1,
                    "sheet_name": sheet_name,
                    "department": sheet_name,
                    "source_sheet_name": f"{os.path.basename(file_path)} [{sheet_name}]",
                    "candidate_name": s_name,
                }

                # Extract all mapped fields (never stop if any optional column is missing)
                for field in [
                    "applied_designation", "negotiated_salary", "recent_designation",
                    "highest_education", "total_experience", "core_technical_skills",
                    "key_responsibilities", "contact_location"
                ]:
                    col_idx = current_col_indices.get(field)
                    val = row[col_idx] if col_idx is not None and col_idx < len(row) else None
                    cand_data[field] = clean_text_value(val)

                # Mobile extraction
                mob_col = current_col_indices.get("mobile")
                raw_mob = row[mob_col] if mob_col is not None and mob_col < len(row) else None
                cand_data["mobile"] = clean_mobile_value(raw_mob)

                # Email extraction
                email_col = current_col_indices.get("email")
                raw_email = row[email_col] if email_col is not None and email_col < len(row) else None
                cand_data["email"] = clean_text_value(raw_email)

                # A valid candidate must have a reasonable name and either a 10-digit mobile or an email
                has_mob = bool(cand_data["mobile"] and len(cand_data["mobile"]) == 10 and cand_data["mobile"].isdigit())
                has_email = bool(cand_data["email"] and "@" in cand_data["email"])

                if (has_mob or has_email) and len(s_name) >= 2:
                    valid_rows.append(cand_data)

    # Strategy 2: CSV / TSV Fallback Extraction
    else:
        try:
            df = read_sheet(file_path)
            sheets_processed.append(os.path.basename(file_path))
            raw_rows = [list(df.columns)] + df.values.tolist()
            current_col_indices = {}

            for r_idx, row in enumerate(raw_rows):
                tok_row = [_tok(c) for c in row]
                matches = 0
                temp_map = {}
                for col_idx, t in enumerate(tok_row):
                    if not t:
                        continue
                    for field, keywords in CANDIDATE_HEADER_SYNONYMS.items():
                        if field not in temp_map:
                            if any(t == kw or t.startswith(kw) or kw.startswith(t) for kw in keywords if len(t) >= 4):
                                temp_map[field] = col_idx
                                matches += 1
                                break

                if "candidate_name" in temp_map and ("mobile" in temp_map or "email" in temp_map or matches >= 3):
                    current_col_indices = temp_map
                    continue

                if not current_col_indices or "candidate_name" not in current_col_indices:
                    continue

                name_col = current_col_indices.get("candidate_name")
                raw_name = row[name_col] if name_col is not None and name_col < len(row) else None
                s_name = clean_text_value(raw_name)

                if not s_name or any(b in s_name.lower() for b in BANNER_PHRASES):
                    continue

                cand_data = {
                    "row_index": r_idx + 1,
                    "sheet_name": os.path.basename(file_path),
                    "department": "",
                    "source_sheet_name": os.path.basename(file_path),
                    "candidate_name": s_name,
                }
                for field in [
                    "applied_designation", "negotiated_salary", "recent_designation",
                    "highest_education", "total_experience", "core_technical_skills",
                    "key_responsibilities", "contact_location"
                ]:
                    col_idx = current_col_indices.get(field)
                    val = row[col_idx] if col_idx is not None and col_idx < len(row) else None
                    cand_data[field] = clean_text_value(val)

                mob_col = current_col_indices.get("mobile")
                raw_mob = row[mob_col] if mob_col is not None and mob_col < len(row) else None
                cand_data["mobile"] = clean_mobile_value(raw_mob)

                email_col = current_col_indices.get("email")
                raw_email = row[email_col] if email_col is not None and email_col < len(row) else None
                cand_data["email"] = clean_text_value(raw_email)

                has_mob = bool(cand_data["mobile"] and len(cand_data["mobile"]) == 10 and cand_data["mobile"].isdigit())
                has_email = bool(cand_data["email"] and "@" in cand_data["email"])
                if (has_mob or has_email) and len(s_name) >= 2:
                    valid_rows.append(cand_data)
        except Exception as ex:
            return {
                "success": False,
                "error": "We couldn't read this CSV file. Please make sure it contains candidate data.",
                "total_sheet_rows": 0,
                "valid_rows": [],
                "sheets_processed": [],
                "errors": [str(ex)],
            }

    if not valid_rows:
        return {
            "success": False,
            "error": "We could not find any candidate records with names and contact details in this spreadsheet. Please check the sheet headers.",
            "total_sheet_rows": 0,
            "valid_rows": [],
            "sheets_processed": sheets_processed,
            "errors": errors,
        }

    return {
        "success": True,
        "error": None,
        "total_sheet_rows": len(valid_rows),
        "valid_rows": valid_rows,
        "sheets_processed": sheets_processed,
        "errors": errors,
    }


def detect_duplicates(
    uploaded_rows: list[dict],
    existing_candidates: list[dict] = None,
    fuzzy_threshold: float = 0.85,
) -> dict:
    """
    Compares uploaded candidate rows against existing candidates in the database.
    A row is flagged as a duplicate if:
      (a) Normalized 10-digit mobile matches an existing candidate's mobile, OR
      (b) Candidate name closely matches an existing candidate (similarity >= fuzzy_threshold)

    Returns structured results:
      {
          "total_rows": int,
          "duplicate_count": int,
          "new_count": int,
          "duplicate_percentage": float,
          "duplicates": list[dict],
          "new_rows": list[dict]
      }
    """
    if existing_candidates is None:
        existing_candidates = db_manager.get_all_candidates() or []

    # Build mobile lookup index for O(1) matching
    mobile_index = {}
    for cand in existing_candidates:
        mob = normalize_mobile(cand.get("mobile", ""))
        if mob and len(mob) >= 10:
            mobile_index[mob] = cand

    duplicates = []
    new_rows = []

    for row in uploaded_rows:
        cand_mobile = normalize_mobile(row.get("mobile", ""))
        cand_name = str(row.get("candidate_name", "")).strip()

        matched_existing = None
        match_type = None
        match_reason = None
        similarity_score = 0.0

        # Rule (a): Exact mobile number match
        if cand_mobile and cand_mobile in mobile_index:
            matched_existing = mobile_index[cand_mobile]
            match_type = "mobile"
            match_reason = f"Mobile number matches existing candidate ({cand_mobile})"
            similarity_score = 1.0

        # Rule (b): Fuzzy name match if mobile didn't match
        if not matched_existing and cand_name:
            best_cand = None
            best_sim = 0.0

            for cand in existing_candidates:
                exist_name = cand.get("candidate_name") or ""
                sim = compute_name_similarity(cand_name, exist_name)
                if sim > best_sim:
                    best_sim = sim
                    best_cand = cand

            if best_sim >= fuzzy_threshold and best_cand:
                matched_existing = best_cand
                match_type = "name"
                pct = int(round(best_sim * 100))
                match_reason = f"Candidate name closely matches '{best_cand.get('candidate_name')}' ({pct}% similarity)"
                similarity_score = best_sim

        if matched_existing:
            duplicates.append({
                "existing_record": matched_existing,
                "uploaded_row": row,
                "match_type": match_type,
                "match_reason": match_reason,
                "similarity": similarity_score,
                "action": "keep",  # Safe default: do NOT overwrite unless HR explicitly chooses
            })
        else:
            new_rows.append(row)

    total = len(uploaded_rows)
    dup_count = len(duplicates)
    new_count = len(new_rows)
    pct = round((dup_count / total * 100), 1) if total > 0 else 0.0

    return {
        "total_rows": total,
        "duplicate_count": dup_count,
        "new_count": new_count,
        "duplicate_percentage": pct,
        "duplicates": duplicates,
        "new_rows": new_rows,
    }


def safe_update_candidate(candidate_id: int, new_data: dict, source_sheet_name: str = "") -> bool:
    """
    Safely updates an existing candidate's record with new uploaded data.
    IMPORTANT: Blank, None, or 'nan' values in new_data will NEVER overwrite
    existing non-empty values in the database.
    Lifecycle fields (status, date_joined, employee_number, relieving_date, etc.)
    are left untouched.
    """
    existing = db_manager.get_candidate_by_id(candidate_id)
    if not existing:
        return False

    updateable_fields = [
        "candidate_name",
        "email",
        "applied_designation",
        "negotiated_salary",
        "recent_designation",
        "highest_education",
        "total_experience",
        "core_technical_skills",
        "key_responsibilities",
        "contact_location",
        "department",
    ]

    updates = []
    params = []

    for field in updateable_fields:
        if field in new_data:
            val = new_data.get(field)
            if val is not None:
                s_val = str(val).strip()
                if s_val and s_val.lower() not in ["nan", "none", "null"]:
                    updates.append(f"{field} = ?")
                    params.append(s_val)

    if source_sheet_name:
        updates.append("source_sheet_name = ?")
        params.append(str(source_sheet_name).strip())

    now = datetime.datetime.now().strftime("%Y-%m-%d %H:%M:%S")
    updates.append("last_updated = ?")
    params.append(now)

    if not updates:
        return True

    params.append(candidate_id)
    conn = db_manager.get_connection()
    cur = conn.cursor()
    try:
        cur.execute(f"UPDATE candidates SET {', '.join(updates)} WHERE id = ?", params)
        conn.commit()
        return True
    except Exception as e:
        print("Safe update error:", e)
        return False
    finally:
        conn.close()


def insert_new_candidate(data: dict, source_sheet_name: str = "") -> bool:
    """
    Inserts a genuinely new candidate record into the candidates table.
    Default status is 'Applied'.
    If a record with the same mobile already exists (e.g. interviewed twice across batches),
    it safely updates and merges non-empty fields into the existing record.
    """
    now = datetime.datetime.now().strftime("%Y-%m-%d %H:%M:%S")
    conn = db_manager.get_connection()
    cur = conn.cursor()
    try:
        cand_mob = str(data.get("mobile") or "").strip() or None
        sheet_src = str(source_sheet_name or data.get("source_sheet_name") or "").strip()

        cur.execute("""
            INSERT INTO candidates
            (candidate_name, mobile, email, applied_designation, negotiated_salary,
             recent_designation, highest_education, total_experience, core_technical_skills,
             key_responsibilities, contact_location, department, status, source_sheet_name, upload_date, last_updated)
            VALUES (?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?)
            ON CONFLICT(mobile) DO UPDATE SET
                candidate_name=COALESCE(NULLIF(excluded.candidate_name, ''), candidates.candidate_name),
                email=COALESCE(NULLIF(excluded.email, ''), candidates.email),
                applied_designation=COALESCE(NULLIF(excluded.applied_designation, ''), candidates.applied_designation),
                negotiated_salary=COALESCE(NULLIF(excluded.negotiated_salary, ''), candidates.negotiated_salary),
                recent_designation=COALESCE(NULLIF(excluded.recent_designation, ''), candidates.recent_designation),
                highest_education=COALESCE(NULLIF(excluded.highest_education, ''), candidates.highest_education),
                total_experience=COALESCE(NULLIF(excluded.total_experience, ''), candidates.total_experience),
                core_technical_skills=COALESCE(NULLIF(excluded.core_technical_skills, ''), candidates.core_technical_skills),
                key_responsibilities=COALESCE(NULLIF(excluded.key_responsibilities, ''), candidates.key_responsibilities),
                contact_location=COALESCE(NULLIF(excluded.contact_location, ''), candidates.contact_location),
                department=COALESCE(NULLIF(excluded.department, ''), candidates.department),
                source_sheet_name=COALESCE(NULLIF(excluded.source_sheet_name, ''), candidates.source_sheet_name),
                last_updated=excluded.last_updated
        """, (
            str(data.get("candidate_name") or "").strip(),
            cand_mob,
            str(data.get("email") or "").strip(),
            str(data.get("applied_designation") or "").strip(),
            str(data.get("negotiated_salary") or "").strip(),
            str(data.get("recent_designation") or "").strip(),
            str(data.get("highest_education") or "").strip(),
            str(data.get("total_experience") or "").strip(),
            str(data.get("core_technical_skills") or "").strip(),
            str(data.get("key_responsibilities") or "").strip(),
            str(data.get("contact_location") or "").strip(),
            str(data.get("department") or "").strip(),
            "Applied",
            sheet_src,
            now,
            now
        ))
        conn.commit()
        return True
    except Exception as e:
        print("Insert candidate error:", e)
        return False
    finally:
        conn.close()


def apply_joining_decisions(
    duplicates: list[dict],
    new_rows: list[dict],
    source_sheet_name: str = "",
) -> dict:
    """
    Applies confirmed decisions safely to the database:
      - 'keep': Candidate record is untouched (left as-is).
      - 'replace': Candidate record updated with non-empty uploaded fields.
      - new_rows: Inserted fresh into candidates table.

    Returns summary:
      {
          "success": bool,
          "added_count": int,
          "updated_count": int,
          "kept_count": int,
          "total_processed": int,
          "errors": list[str]
      }
    """
    added_count = 0
    updated_count = 0
    kept_count = 0
    errors = []

    # 1. Process duplicates based on user action
    for item in duplicates:
        action = item.get("action", "keep")
        cand_id = item.get("existing_record", {}).get("id")
        uploaded = item.get("uploaded_row", {})
        cand_name = uploaded.get("candidate_name") or item.get("existing_record", {}).get("candidate_name") or "Candidate"
        sheet_source = source_sheet_name or uploaded.get("source_sheet_name") or ""

        if action == "replace":
            if cand_id:
                ok = safe_update_candidate(cand_id, uploaded, source_sheet_name=sheet_source)
                if ok:
                    updated_count += 1
                else:
                    errors.append(f"Could not update record for '{cand_name}'.")
            else:
                errors.append(f"Missing ID for existing candidate '{cand_name}'.")
        else:
            # Action is 'keep' -> untouched!
            kept_count += 1

    # 2. Process genuinely new candidates
    for new_cand in new_rows:
        cand_name = new_cand.get("candidate_name") or "New Candidate"
        sheet_source = source_sheet_name or new_cand.get("source_sheet_name") or ""
        ok = insert_new_candidate(new_cand, source_sheet_name=sheet_source)
        if ok:
            added_count += 1
        else:
            errors.append(f"Could not add new candidate '{cand_name}'.")

    return {
        "success": len(errors) == 0,
        "added_count": added_count,
        "updated_count": updated_count,
        "kept_count": kept_count,
        "total_processed": added_count + updated_count + kept_count,
        "errors": errors,
    }


def process_sheet(file_path: str, source_sheet_name: str):
    """
    Legacy wrapper for backwards-compatibility with older tests.
    Extracts rows, detects duplicates, safely updates/adds.
    """
    extracted = extract_candidate_rows_from_sheet(file_path)
    if not extracted["success"]:
        return {"success": False, "error": extracted["error"]}

    valid_rows = extracted["valid_rows"]
    dup_res = detect_duplicates(valid_rows)

    # In legacy non-interactive mode, new rows are inserted and duplicates are kept
    res = apply_joining_decisions(
        duplicates=dup_res["duplicates"],
        new_rows=dup_res["new_rows"],
        source_sheet_name=source_sheet_name
    )

    return {
        "success": True,
        "total_rows": extracted["total_sheet_rows"],
        "processed": res["added_count"] + res["updated_count"],
        "duplicates_skipped": res["kept_count"],
        "errors": extracted["errors"] + res["errors"],
    }
