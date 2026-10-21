"""
Service: placeholder_engine.py
Purpose: Robust placeholder replacement engine for python-docx templates.
Handles single ({Key}) and double ({{Key}}) brace syntax, cross-run splits,
multi-line line breaks (WD_BREAK.LINE), document tables, headers, and footers.
"""

from typing import Dict, Any
from docx.enum.text import WD_BREAK


def _replace_token_in_paragraph(p, search_token: str, replace_value: str) -> None:
    """
    Replaces all occurrences of `search_token` in paragraph `p`, even if
    `search_token` is split across multiple runs by Microsoft Word.
    Preserves font styling of the starting run. Converts any '\\n' in
    `replace_value` into native Word soft line-breaks (WD_BREAK.LINE).
    """
    if not p.runs or not search_token:
        return

    while True:
        full_text = "".join(r.text for r in p.runs)
        idx = full_text.find(search_token)
        if idx == -1:
            break

        end_idx = idx + len(search_token)

        pos = 0
        first_overlap = True

        for r in p.runs:
            r_len = len(r.text)
            r_start = pos
            r_end = pos + r_len
            pos += r_len

            # Check if this run overlaps with [idx, end_idx)
            if r_end <= idx or r_start >= end_idx:
                continue

            prefix = r.text[:max(0, idx - r_start)]
            suffix = r.text[max(0, end_idx - r_start):]

            if first_overlap:
                first_overlap = False
                # If replacement value contains newlines (e.g. multi-line address),
                # insert native Word line-breaks into this run
                if "\n" in replace_value:
                    lines = replace_value.split("\n")
                    r.text = prefix + lines[0]
                    for line in lines[1:]:
                        r.add_break(WD_BREAK.LINE)
                        r.add_text(line)
                    if suffix:
                        r.add_text(suffix)
                else:
                    r.text = prefix + replace_value + suffix
            else:
                # Subsequent runs spanned by the placeholder are cleared of the matched text
                r.text = suffix


def replace_placeholders_in_docx(doc, mapping: Dict[str, Any]) -> None:
    """
    Replaces all placeholders in a loaded python-docx Document.

    Arguments:
        doc: Already-loaded docx.Document object.
        mapping: Dictionary of clean keys to values, e.g.:
                 {"Full Name": "Rahul Sharma", "Job Title": "Field Coordinator", ...}
                 Keys can be provided with or without curly braces.

    Features:
        - Replaces BOTH '{{Key}}' and '{Key}' formats.
        - Resolves placeholders split across multiple Word runs.
        - Replaces '\\n' in values with native Word soft line breaks.
        - Processes document body paragraphs, all table cells, and headers/footers.
    """
    if not doc or not mapping:
        return

    # Build token list with both double and single brace variants
    tokens = []
    for k, v in mapping.items():
        clean_k = str(k).strip()
        # Strip enclosing braces if caller provided them
        if clean_k.startswith("{{") and clean_k.endswith("}}"):
            clean_k = clean_k[2:-2].strip()
        elif clean_k.startswith("{") and clean_k.endswith("}"):
            clean_k = clean_k[1:-1].strip()

        val_str = str(v) if v is not None else ""

        # Double curly brace variant (highest precedence)
        tokens.append((f"{{{{{clean_k}}}}}", val_str))
        # Single curly brace variant
        tokens.append((f"{{{clean_k}}}", val_str))

    # Sort tokens by length in descending order so {{Key}} matches before {Key}
    tokens.sort(key=lambda item: len(item[0]), reverse=True)

    def process_paragraphs(paragraphs):
        for p in paragraphs:
            # Quick check if any token could be in this paragraph
            p_text = "".join(r.text for r in p.runs)
            if not p_text:
                continue
            for token, val in tokens:
                if token in p_text:
                    _replace_token_in_paragraph(p, token, val)
                    # Refresh paragraph text after replacement for subsequent tokens
                    p_text = "".join(r.text for r in p.runs)

    # 1. Process all main body paragraphs
    process_paragraphs(doc.paragraphs)

    # 2. Process all tables and nested cell paragraphs
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                process_paragraphs(cell.paragraphs)

    # 3. Process headers and footers across all document sections
    for section in doc.sections:
        if section.header:
            process_paragraphs(section.header.paragraphs)
        if section.footer:
            process_paragraphs(section.footer.paragraphs)
