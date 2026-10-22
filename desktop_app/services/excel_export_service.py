"""
Service: excel_export_service.py
Purpose: Professional executive-grade Excel (.xlsx) export utility with Doctors For You branding.
Features:
- DFY Dark Navy (#0A2540) header banner with organization title and report metadata
- Deep Corporate Blue (#1E3A8A) table headers with auto-filter and freeze panes (A6)
- Alternating zebra striping (#FFFFFF and #F8FAFC)
- Smart column formatting (Indian Currency ₹ #,##0, DD-MMM-YYYY dates, center/left alignments)
- Content-aware column auto-fitting (min 14, max 55)
- Multi-sheet workbook support with customized per-sheet banners & subtitles
"""

import os
from datetime import datetime, date
from typing import List, Dict, Any, Optional

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter


def _try_parse_currency(val: Any) -> Any:
    """Attempts to parse a value into a numeric currency amount."""
    if val is None or val == "":
        return None
    if isinstance(val, (int, float)):
        return val
    try:
        cleaned = str(val).replace("₹", "").replace(",", "").strip()
        if not cleaned:
            return None
        if "." in cleaned:
            return float(cleaned)
        return int(cleaned)
    except (ValueError, TypeError):
        return val


def _try_parse_date(val: Any) -> Any:
    """Attempts to parse standard date strings into datetime.date objects."""
    if val is None or val == "":
        return None
    if isinstance(val, (datetime, date)):
        return val
    if isinstance(val, str):
        val_str = val.strip()
        for fmt in (
            "%Y-%m-%d",
            "%Y-%m-%d %H:%M:%S",
            "%d-%m-%Y",
            "%d/%m/%Y",
            "%d-%b-%Y",
            "%Y/%m/%d",
            "%d.%m.%Y",
        ):
            try:
                dt = datetime.strptime(val_str, fmt)
                return dt.date()
            except ValueError:
                continue
    return val


def _get_column_type(col_name: str) -> str:
    """Classifies column name to determine formatting and alignment."""
    name = str(col_name).strip().lower()
    currency_keywords = (
        "salary", "fee", "amount", "stipend", "ctc", "budget",
        "cost", "pay", "gross", "net", "allowance", "compensation"
    )
    date_keywords = (
        "date", "dob", "doj", "joined", "created_at", "updated_at", "timestamp"
    )
    center_keywords = (
        "mobile", "phone", "contact", "status", "id", "code",
        "gender", "blood group", "state", "pincode", "pin code",
        "pan", "aadhaar", "uan", "emp id", "priority", "urgent", "type"
    )

    if any(k in name for k in currency_keywords):
        return "currency"
    if any(k in name for k in date_keywords):
        return "date"
    if any(k in name for k in center_keywords):
        return "center"
    return "text"


def populate_and_style_worksheet(
    ws,
    headers: List[str],
    rows: List[List[Any]],
    title: str = "DOCTORS FOR YOU — HR & RECRUITMENT OPERATIONS",
    subtitle: str = ""
) -> None:
    """
    Populates and styles an individual openpyxl Worksheet with executive DFY branding.

    Arguments:
        ws: openpyxl Worksheet object
        headers: List of column header names
        rows: List of row data lists
        title: Header banner text
        subtitle: Subtitle / report metadata text
    """
    if ws.views.sheetView:
        ws.views.sheetView[0].showGridLines = True

    num_cols = max(1, len(headers))
    last_col_letter = get_column_letter(num_cols)

    # ==========================================
    # 1. DFY Organization Header Banner (Rows 1-2)
    # ==========================================
    banner_fill = PatternFill(start_color="FF0A2540", end_color="FF0A2540", fill_type="solid")
    banner_font = Font(name="Calibri", size=14, bold=True, color="FFFFFFFF")
    banner_align = Alignment(horizontal="left", vertical="center", indent=1)

    ws.merge_cells(start_row=1, start_column=1, end_row=2, end_column=num_cols)
    for r in range(1, 3):
        ws.row_dimensions[r].height = 20
        for c in range(1, num_cols + 1):
            cell = ws.cell(row=r, column=c)
            cell.fill = banner_fill

    title_cell = ws.cell(row=1, column=1)
    title_cell.value = title or "DOCTORS FOR YOU — HR & RECRUITMENT OPERATIONS"
    title_cell.font = banner_font
    title_cell.alignment = banner_align

    # ==========================================
    # 2. Report Subtitle / Metadata (Row 3)
    # ==========================================
    now_str = datetime.now().strftime("%d-%b-%Y %I:%M %p")
    if subtitle:
        if "Generated:" not in subtitle:
            subtitle_text = f"{subtitle} | Generated: {now_str} | Confidential"
        else:
            subtitle_text = subtitle
    else:
        sheet_title = getattr(ws, "title", "Report")
        subtitle_text = f"Report: {sheet_title} | Generated: {now_str} | Confidential"

    ws.row_dimensions[3].height = 18
    if num_cols > 1:
        ws.merge_cells(start_row=3, start_column=1, end_row=3, end_column=num_cols)

    sub_font = Font(name="Calibri", size=10, italic=True, color="FF475569")
    sub_align = Alignment(horizontal="left", vertical="center", indent=1)
    sub_cell = ws.cell(row=3, column=1)
    sub_cell.value = subtitle_text
    sub_cell.font = sub_font
    sub_cell.alignment = sub_align

    # ==========================================
    # 3. Separator Row (Row 4)
    # ==========================================
    ws.row_dimensions[4].height = 10

    # ==========================================
    # 4. Table Header Row (Row 5)
    # ==========================================
    ws.row_dimensions[5].height = 26
    header_fill = PatternFill(start_color="FF1E3A8A", end_color="FF1E3A8A", fill_type="solid")
    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFFFF")
    header_align = Alignment(horizontal="center", vertical="center", wrap_text=True)
    header_border = Border(
        left=Side(style="thin", color="FFCBD5E1"),
        right=Side(style="thin", color="FFCBD5E1"),
        top=Side(style="thin", color="FFCBD5E1"),
        bottom=Side(style="medium", color="FFCBD5E1"),
    )

    col_types = [_get_column_type(h) for h in headers]

    for col_idx, col_name in enumerate(headers, start=1):
        cell = ws.cell(row=5, column=col_idx)
        cell.value = col_name
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = header_align
        cell.border = header_border

    # Enable Auto-Filter and Freeze Panes below table header
    ws.auto_filter.ref = f"A5:{last_col_letter}5"
    ws.freeze_panes = "A6"

    # ==========================================
    # 5. Data Rows (Row 6 onwards)
    # ==========================================
    data_font = Font(name="Calibri", size=10)
    zebra_white = PatternFill(start_color="FFFFFFFF", end_color="FFFFFFFF", fill_type="solid")
    zebra_gray = PatternFill(start_color="FFF8FAFC", end_color="FFF8FAFC", fill_type="solid")
    cell_border = Border(
        left=Side(style="thin", color="FFE2E8F0"),
        right=Side(style="thin", color="FFE2E8F0"),
        top=Side(style="thin", color="FFE2E8F0"),
        bottom=Side(style="thin", color="FFE2E8F0"),
    )

    for row_offset, row_data in enumerate(rows):
        current_row = 6 + row_offset
        ws.row_dimensions[current_row].height = 20
        # Alternating striping: Even row index white, odd row index light slate gray
        row_fill = zebra_gray if (current_row % 2 == 1) else zebra_white

        for col_idx, val in enumerate(row_data, start=1):
            cell = ws.cell(row=current_row, column=col_idx)
            cell.font = data_font
            cell.fill = row_fill
            cell.border = cell_border

            ctype = col_types[col_idx - 1] if col_idx - 1 < len(col_types) else "text"

            if ctype == "currency":
                parsed_cur = _try_parse_currency(val)
                if isinstance(parsed_cur, (int, float)):
                    cell.value = parsed_cur
                    cell.number_format = '"₹" #,##0'
                else:
                    cell.value = val if val is not None else ""
                cell.alignment = Alignment(horizontal="right", vertical="center")

            elif ctype == "date":
                parsed_dt = _try_parse_date(val)
                if isinstance(parsed_dt, (date, datetime)):
                    cell.value = parsed_dt
                    cell.number_format = "DD-MMM-YYYY"
                else:
                    cell.value = val if val is not None else ""
                    cell.number_format = "DD-MMM-YYYY"
                cell.alignment = Alignment(horizontal="center", vertical="center")

            elif ctype == "center":
                cell.value = val if val is not None else ""
                cell.alignment = Alignment(horizontal="center", vertical="center")

            else:
                cell.value = val if val is not None else ""
                cell.alignment = Alignment(horizontal="left", vertical="center", indent=1)

    # ==========================================
    # 6. Column Width Auto-Fitting
    # ==========================================
    max_data_row = max(5, 5 + len(rows))
    for col_idx in range(1, num_cols + 1):
        col_letter = get_column_letter(col_idx)
        header_val = str(ws.cell(row=5, column=col_idx).value or "")
        max_len = len(header_val)

        for r in range(6, max_data_row + 1):
            c_val = ws.cell(row=r, column=col_idx).value
            if c_val is not None:
                if isinstance(c_val, (date, datetime)):
                    val_str = c_val.strftime("%d-%b-%Y")
                elif isinstance(c_val, (int, float)) and col_idx - 1 < len(col_types) and col_types[col_idx - 1] == "currency":
                    val_str = f"₹ {c_val:,.0f}"
                else:
                    val_str = str(c_val)
                max_len = max(max_len, len(val_str))

        calculated_width = max(max_len + 4, 14)
        ws.column_dimensions[col_letter].width = min(calculated_width, 55)


def export_to_excel(data: List[Dict[str, Any]], output_path: str, sheet_name: str = "Report") -> str:
    """
    Exports a list of dictionaries to an executive-styled Excel (.xlsx) workbook.

    Arguments:
        data: List of dictionaries representing table rows.
        output_path: Absolute or relative destination file path.
        sheet_name: Name of the worksheet tab.

    Returns:
        output_path on success.
    """
    parent_dir = os.path.dirname(os.path.abspath(output_path))
    if parent_dir:
        os.makedirs(parent_dir, exist_ok=True)

    df = pd.DataFrame(data if data else [])
    df = df.where(pd.notnull(df), None)

    if df.empty and not df.columns.tolist():
        headers = ["Status"]
        rows = [["No records to display"]]
    else:
        headers = [str(c) for c in df.columns.tolist()]
        rows = df.values.tolist()

    wb = Workbook()
    ws = wb.active
    ws.title = sheet_name[:31]

    populate_and_style_worksheet(
        ws=ws,
        headers=headers,
        rows=rows,
        title="DOCTORS FOR YOU — HR & RECRUITMENT OPERATIONS",
        subtitle=f"Report: {sheet_name}"
    )

    wb.save(output_path)
    wb.close()
    return output_path


def export_multisheet_to_excel(
    sheets_data: Dict[str, Any],
    output_path: str,
    main_title: str = "DOCTORS FOR YOU — HR & RECRUITMENT OPERATIONS",
    sheet_subtitles: Optional[Dict[str, str]] = None
) -> str:
    """
    Exports multiple sheets of data into a single executive-styled Excel workbook.

    Arguments:
        sheets_data: Dictionary mapping sheet names to data. Data can be:
                     - List of dicts: [{"col1": val, ...}, ...]
                     - Dict with {"data": [...], "subtitle": "...", "title": "..."}
                     - pandas DataFrame
                     - Dict with {"headers": [...], "rows": [...]}
        output_path: Destination file path.
        main_title: Organization header banner title.
        sheet_subtitles: Optional mapping of sheet_name -> subtitle string.

    Returns:
        output_path on success.
    """
    parent_dir = os.path.dirname(os.path.abspath(output_path))
    if parent_dir:
        os.makedirs(parent_dir, exist_ok=True)

    wb = Workbook()
    first_sheet = True

    for sheet_name, sheet_content in sheets_data.items():
        clean_sheet_name = str(sheet_name)[:31]
        if first_sheet:
            ws = wb.active
            ws.title = clean_sheet_name
            first_sheet = False
        else:
            ws = wb.create_sheet(title=clean_sheet_name)

        title = main_title
        subtitle = (sheet_subtitles or {}).get(sheet_name, "")

        data = sheet_content
        if isinstance(sheet_content, dict):
            if "headers" in sheet_content and "rows" in sheet_content:
                headers = [str(h) for h in sheet_content["headers"]]
                rows = sheet_content["rows"]
                title = sheet_content.get("title", title)
                subtitle = sheet_content.get("subtitle", subtitle)
                populate_and_style_worksheet(ws, headers, rows, title=title, subtitle=subtitle)
                continue
            elif "data" in sheet_content:
                data = sheet_content["data"]
                title = sheet_content.get("title", title)
                subtitle = sheet_content.get("subtitle", subtitle)

        if isinstance(data, pd.DataFrame):
            df = data.where(pd.notnull(data), None)
        else:
            df = pd.DataFrame(data if data else [])
            df = df.where(pd.notnull(df), None)

        if df.empty and not df.columns.tolist():
            headers = ["Status"]
            rows = [["No records to display"]]
        else:
            headers = [str(c) for c in df.columns.tolist()]
            rows = df.values.tolist()

        populate_and_style_worksheet(ws, headers, rows, title=title, subtitle=subtitle)

    wb.save(output_path)
    wb.close()
    return output_path
