import os
import sys
import tempfile
import openpyxl

SCRIPT_DIR = os.path.dirname(os.path.abspath(__file__))
DESKTOP_APP_DIR = os.path.dirname(SCRIPT_DIR)
if DESKTOP_APP_DIR not in sys.path:
    sys.path.insert(0, DESKTOP_APP_DIR)
sys.stdout.reconfigure(encoding="utf-8")

from services import excel_export_service
from services import hr_report_exporter


def test_excel_export_styling():
    print("Testing Excel Export Styling...")
    sample_data = [
        {"Candidate Name": "Aarav Patel", "Mobile": "9876543210", "Salary": 35000, "Date Joined": "2026-11-01", "Status": "Selected"},
        {"Candidate Name": "Diya Sharma", "Mobile": "9876543211", "Salary": 42000, "Date Joined": "2026-11-05", "Status": "Selected"},
        {"Candidate Name": "Kabir Khan", "Mobile": "9876543212", "Salary": 28000, "Date Joined": "2026-11-10", "Status": "Applied"},
    ]

    with tempfile.TemporaryDirectory() as tmp_dir:
        out_path = os.path.join(tmp_dir, "test_styled_export.xlsx")
        excel_export_service.export_to_excel(sample_data, out_path, sheet_name="Candidates")

        assert os.path.exists(out_path), "Excel file was not created"
        wb = openpyxl.load_workbook(out_path)
        ws = wb["Candidates"]

        # 1. Header Banner Check
        assert ws.cell(row=1, column=1).value is not None, "Missing title banner in row 1"
        assert "DOCTORS FOR YOU" in str(ws.cell(row=1, column=1).value), "Banner missing DOCTORS FOR YOU"

        # 2. Table Headers Check (Row 5)
        header_vals = [ws.cell(row=5, column=col).value for col in range(1, 6)]
        assert "Candidate Name" in header_vals
        assert "Salary" in header_vals

        # 3. Freeze panes check
        assert ws.freeze_panes in ["A6", "A5"], f"Unexpected freeze panes: {ws.freeze_panes}"

        # 4. Auto-filter check
        assert ws.auto_filter.ref is not None, "Auto-filter not set on table"

        # 5. Styling Check (Row 5 header background)
        header_cell = ws.cell(row=5, column=1)
        assert header_cell.fill.fill_type == "solid", "Header cell should have solid fill"
        assert header_cell.fill.start_color.rgb in ["FF1E3A8A", "1E3A8A", "FF0A2540", "0A2540"], f"Unexpected header color: {header_cell.fill.start_color.rgb}"

        # 6. Zebra striping check
        row6_cell = ws.cell(row=6, column=1)
        row7_cell = ws.cell(row=7, column=1)
        # One of row 6 / row 7 should have the alternating background
        assert row7_cell.fill.fill_type == "solid", "Row 7 cell should have zebra fill"
        assert row7_cell.fill.start_color.rgb in ["FFF8FAFC", "F8FAFC"], f"Unexpected zebra color: {row7_cell.fill.start_color.rgb}"

        # 7. Number format check (Salary & Date)
        salary_col_idx = header_vals.index("Salary") + 1
        date_col_idx = header_vals.index("Date Joined") + 1
        salary_cell = ws.cell(row=6, column=salary_col_idx)
        date_cell = ws.cell(row=6, column=date_col_idx)

        assert "₹" in str(salary_cell.number_format) or "#,##0" in str(salary_cell.number_format), f"Salary number format invalid: {salary_cell.number_format}"
        assert "YY" in str(date_cell.number_format).upper() or "MMM" in str(date_cell.number_format).upper(), f"Date format invalid: {date_cell.number_format}"

        # 8. Empty data check
        empty_out = os.path.join(tmp_dir, "test_empty.xlsx")
        excel_export_service.export_to_excel([], empty_out, sheet_name="Empty")
        assert os.path.exists(empty_out), "Empty export failed to create file"

        print("  ✓ All Excel styling verifications passed!")


def test_daily_briefing_styling():
    print("Testing Daily Briefing 2-Sheet Excel Styling...")
    with tempfile.TemporaryDirectory() as tmp_dir:
        out_path = os.path.join(tmp_dir, "test_daily_briefing.xlsx")
        result_path = hr_report_exporter.export_daily_executive_summary(target_path=out_path)

        assert os.path.exists(result_path), "Daily briefing Excel file was not created"
        wb = openpyxl.load_workbook(result_path)

        expected_sheets = ["Executive Overview", "Action Priorities"]
        for sheet_name in expected_sheets:
            assert sheet_name in wb.sheetnames, f"Missing expected sheet: {sheet_name}"
            ws = wb[sheet_name]

            # 1. Row 1 banner containing DOCTORS FOR YOU
            banner_val = str(ws.cell(row=1, column=1).value or "")
            assert "DOCTORS FOR YOU" in banner_val, f"{sheet_name}: Row 1 banner missing DOCTORS FOR YOU (got: {banner_val})"

            # 2. Subtitle check
            sub_val = str(ws.cell(row=3, column=1).value or "")
            if sheet_name == "Executive Overview":
                assert "Daily Operations & Pipeline KPI Summary" in sub_val, f"Executive Overview subtitle missing: {sub_val}"
            elif sheet_name == "Action Priorities":
                assert "Daily Operational Action Items & Directives" in sub_val, f"Action Priorities subtitle missing: {sub_val}"

            # 3. Freeze panes set
            assert ws.freeze_panes in ["A6", "A5"], f"{sheet_name}: Unexpected freeze panes: {ws.freeze_panes}"

            # 4. Auto-filter set
            assert ws.auto_filter.ref is not None, f"{sheet_name}: Auto-filter not set on table"

            # 5. Solid fill header row
            header_cell = ws.cell(row=5, column=1)
            assert header_cell.fill.fill_type == "solid", f"{sheet_name}: Header cell should have solid fill"
            assert header_cell.fill.start_color.rgb in ["FF1E3A8A", "1E3A8A", "FF0A2540", "0A2540"], f"{sheet_name}: Unexpected header color: {header_cell.fill.start_color.rgb}"

            # 6. Data rows present
            assert ws.cell(row=6, column=1).value is not None, f"{sheet_name}: Row 6 should contain data"

    print("  ✓ All Daily Briefing styling verifications passed!")


if __name__ == "__main__":
    test_excel_export_styling()
    test_daily_briefing_styling()
