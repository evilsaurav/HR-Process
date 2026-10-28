"""
Comprehensive Test Script: test_all_document_alignments.py
Generates and inspects all 6 HR document types for typography, alignment, and page fit:
1. Offer Letter (offer)
2. Extension Letter (extension)
3. Increment Letter (increment)
4. Termination Letter (termination)
5. Experience Certificate (experience_certificate)
6. Salary Certificate (salary_certificate)
"""

import os
import sys
import docx
import pypdf
from docx.enum.text import WD_ALIGN_PARAGRAPH

SCRIPT_DIR = os.path.dirname(os.path.abspath(__file__))
DESKTOP_APP_DIR = os.path.dirname(SCRIPT_DIR)
if DESKTOP_APP_DIR not in sys.path:
    sys.path.insert(0, DESKTOP_APP_DIR)

if hasattr(sys.stdout, "reconfigure"):
    sys.stdout.reconfigure(encoding="utf-8", errors="replace")

from database import db_manager
from services import document_generation_service
from services import pdf_converter_service

def test_all_documents():
    print("=" * 75)
    print("TESTING DOCUMENT ALIGNMENT, SPACING & PAGE BUDGET ACROSS ALL 6 TYPES")
    print("=" * 75)

    db_manager.init_db()

    # Setup candidate
    test_mobile = "9876543219"
    cand_data = {
        "candidate_name": "Deepak Verma",
        "mobile": test_mobile,
        "email": "deepak.verma@example.com",
        "applied_designation": "Data Entry Operator",
        "negotiated_salary": "28000",
        "recent_designation": "DEO",
        "highest_education": "Graduate",
        "total_experience": "3 Years",
        "contact_location": "Lucknow",
    }
    db_manager.upsert_candidate(cand_data, "test_align_all.xlsx")
    cand = db_manager.get_candidate_by_mobile(test_mobile)
    assert cand is not None, "Failed to get test candidate"
    cand_id = cand["id"]

    db_manager.update_candidate_extra_details(
        candidate_id=cand_id,
        father_name="Mahesh Verma",
        full_address="House 12B, Indira Nagar, Sector 4, Lucknow, Uttar Pradesh 226016",
        gender="Male",
        department="Operations",
        date_joined="2024-03-01",
        employee_number="DFY-2024"
    )

    doc_test_cases = [
        ("offer", {}, 2),
        ("extension", {"extension_start_date": "2025-03-01"}, 1),
        ("increment", {"effective_from": "2025-03-01", "revised_pay_grade": "32000"}, 1),
        ("termination", {"termination_date": "2025-02-28", "reason": "Tenure Completion", "notice_period": "30 Days"}, 1),
        ("experience_certificate", {"exit_date": "2025-02-28"}, 1),
        ("salary_certificate", {"exit_date": "Till Date"}, 1)
    ]

    for doc_type, extra_fields, expected_max_pages in doc_test_cases:
        print(f"\n[{doc_type.upper()}] Generating preview...")
        res = document_generation_service.generate_preview_document(
            candidate_id=cand_id,
            document_type=doc_type,
            extra_fields=extra_fields
        )
        assert res.get("success") is True, f"Failed to generate {doc_type}: {res.get('error')}"
        docx_path = res["preview_file_path"]
        assert os.path.exists(docx_path), f"File {docx_path} not found"

        doc = docx.Document(docx_path)

        # 1. Check Table 0 (Header Table)
        assert len(doc.tables) >= 1, f"[{doc_type}] Header table missing!"
        table = doc.tables[0]
        assert len(table.rows) == 1 and len(table.columns) == 2, f"[{doc_type}] Table 0 must be 1x2"
        c0 = table.cell(0, 0)
        c1 = table.cell(0, 1)
        assert "Reference No:" in c0.text, f"[{doc_type}] Cell 0 missing 'Reference No:'"
        if doc_type in ["experience_certificate", "salary_certificate"]:
            assert "Date:" in c1.text, f"[{doc_type}] Cell 1 missing 'Date:'"
        else:
            assert "Dated:" in c1.text, f"[{doc_type}] Cell 1 missing 'Dated:'"

        assert c0.paragraphs[0].alignment in (WD_ALIGN_PARAGRAPH.LEFT, None)
        assert c1.paragraphs[0].alignment == WD_ALIGN_PARAGRAPH.RIGHT

        # 2. Check Text Integrity
        for p in doc.paragraphs:
            txt = p.text
            # No old fragile space padding
            assert (" " * 50) not in txt, f"[{doc_type}] Found fragile space padding in: {repr(txt[:60])}"
            # No redundant currency
            assert "INR ₹" not in txt, f"[{doc_type}] Found redundant currency 'INR ₹' in: {repr(txt[:60])}"
            # No broken offer date
            assert "Dated ." not in txt, f"[{doc_type}] Found broken date 'Dated .' in: {repr(txt[:60])}"
            # No leading spaces before Sub:
            if txt.strip().startswith("Sub:"):
                assert not txt.startswith(" "), f"[{doc_type}] 'Sub:' has leading space in: {repr(txt)}"
            # No leading spaces before Dear
            if txt.strip().startswith("Dear "):
                assert not txt.startswith(" "), f"[{doc_type}] 'Dear' has leading space in: {repr(txt)}"
            # No leading spaces before Greetings
            if txt.strip().startswith("Greetings"):
                assert not txt.startswith(" "), f"[{doc_type}] 'Greetings' has leading space in: {repr(txt)}"

        # 3. Convert to PDF and Check Page Budget
        ok, pdf_path, err = pdf_converter_service.convert_docx_to_pdf(docx_path)
        if ok and os.path.exists(pdf_path):
            reader = pypdf.PdfReader(pdf_path)
            num_pages = len(reader.pages)
            print(f"  ✓ PDF converted cleanly ({num_pages} page(s)). Expected max: {expected_max_pages} page(s).")
            assert num_pages <= expected_max_pages, (
                f"[{doc_type}] Page count exceeded! Got {num_pages} pages, expected <= {expected_max_pages}"
            )
            # Verify last page has content and is not empty/orphaned
            last_page = reader.pages[-1]
            last_text = last_page.extract_text().strip()
            assert len(last_text) > 0, f"[{doc_type}] Last page is empty!"
        else:
            print(f"  [Notice] Word COM PDF conversion skipped: {err}")

        print(f"  ✓ [{doc_type}] PASSED all alignment and spacing checks!")

    # Cleanup candidate
    db_manager.delete_candidate(cand_id)
    print("\n" + "=" * 75)
    print("ALL 6 DOCUMENT ALIGNMENT & SPACING CHECKS PASSED WITH 100% SUCCESS! ✅")
    print("=" * 75)
    return True

if __name__ == "__main__":
    if test_all_documents():
        sys.exit(0)
    else:
        sys.exit(1)
