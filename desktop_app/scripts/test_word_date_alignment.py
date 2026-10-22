"""
Verification Script: test_word_date_alignment.py
Validates Task 3: Word Template Header Date Alignment Fix.

Verifies:
1. Preview generation across all core document types:
   - offer
   - extension
   - increment
   - termination
   - experience_certificate
2. Header table verification for each generated doc:
   - len(doc.tables) >= 1
   - doc.tables[0] has exactly 1 row and 2 columns
   - doc.tables[0].cell(0, 0) contains "Reference No:" and is left-aligned
   - doc.tables[0].cell(0, 1) contains "Dated:" or "Date:" and is right-aligned
   - Reference No and Date are in separate cells, eliminating fragile space padding
3. Formatter verification:
   - format_date_long cleans timestamp fragments (e.g. 00:00:00) and formats cleanly to 'DD MMMM YYYY'
   - calculate_end_date calculates date + 1 year - 1 day cleanly
"""

import os
import sys
import docx
from docx.enum.text import WD_ALIGN_PARAGRAPH

# Ensure desktop_app is in sys.path
SCRIPT_DIR = os.path.dirname(os.path.abspath(__file__))
DESKTOP_APP_DIR = os.path.dirname(SCRIPT_DIR)
if DESKTOP_APP_DIR not in sys.path:
    sys.path.insert(0, DESKTOP_APP_DIR)

sys.stdout.reconfigure(encoding="utf-8")

from database import db_manager
from services import document_generation_service
from services.document_formatters import format_date_long, calculate_end_date


def setup_test_candidate():
    db_manager.init_db()

    # Ensure reference sequence exists
    ref_cfg = db_manager.get_reference_config("hr_reference")
    if not ref_cfg or ref_cfg["current_number"] + 50 > ref_cfg["range_end"]:
        db_manager.save_reference_config("hr_reference", "DFY/HR/{YEAR}/{NUM}", 3, 1, 9999)

    test_mobile = "9988776655"
    cand_data = {
        "candidate_name": "Priya Sharma",
        "mobile": test_mobile,
        "email": "priya.sharma@example.com",
        "applied_designation": "Data Entry Operator",
        "negotiated_salary": "32000",
        "recent_designation": "DEO",
        "highest_education": "Post Graduate",
        "total_experience": "4 Years",
        "contact_location": "Delhi",
    }
    db_manager.upsert_candidate(cand_data, "test_alignment.xlsx")
    cand = db_manager.get_candidate_by_mobile(test_mobile)
    assert cand is not None, "Failed to upsert test candidate"
    cand_id = cand["id"]

    db_manager.update_candidate_extra_details(
        candidate_id=cand_id,
        father_name="Ramesh Sharma",
        full_address="House 42, Block C, Preet Vihar, Delhi",
        gender="Female",
        department="Health",
        date_joined="2024-05-01",
        employee_number="DFY-9901"
    )
    return cand_id


def test_formatters():
    print("\n--- Testing Date Formatters ---")
    
    # 1. format_date_long with timestamp fragments
    t1 = format_date_long("2026-10-15 00:00:00")
    assert t1 == "15 October 2026", f"Expected '15 October 2026', got '{t1}'"
    
    t2 = format_date_long("2026-10-15T00:00:00")
    assert t2 == "15 October 2026", f"Expected '15 October 2026', got '{t2}'"

    t3 = format_date_long("15-10-2026 00:00:00")
    assert t3 == "15 October 2026", f"Expected '15 October 2026', got '{t3}'"

    # 2. calculate_end_date
    end_date = calculate_end_date("2024-04-01")
    assert end_date == "31 March 2025", f"Expected '31 March 2025', got '{end_date}'"

    end_date_iso = calculate_end_date("2024-04-01 00:00:00")
    assert end_date_iso == "31 March 2025", f"Expected '31 March 2025', got '{end_date_iso}'"

    print("✅ Formatter tests passed successfully.")


def test_document_date_alignment():
    print("\n--- Testing Document Date Alignment Tables ---")
    candidate_id = setup_test_candidate()

    test_cases = [
        ("offer", {}),
        ("extension", {"extension_start_date": "2025-05-01"}),
        ("increment", {"effective_from": "2025-05-01", "revised_pay_grade": "35000"}),
        ("termination", {"termination_date": "2025-04-30", "reason": "Contract Completion", "notice_period": "30 Days"}),
        ("experience_certificate", {"exit_date": "2025-04-30"}),
    ]

    for doc_type, extra_fields in test_cases:
        print(f"\nTesting document type: '{doc_type}'...")
        res = document_generation_service.generate_preview_document(
            candidate_id=candidate_id,
            document_type=doc_type,
            extra_fields=extra_fields
        )
        assert res.get("success") is True, f"Preview generation failed for {doc_type}: {res.get('error')}"
        file_path = res["preview_file_path"]
        assert os.path.exists(file_path), f"File {file_path} not found"

        doc = docx.Document(file_path)

        # Assertion 1: Must have at least 1 table
        assert len(doc.tables) >= 1, f"[{doc_type}] Expected doc.tables to have at least 1 table, found {len(doc.tables)}"

        # Assertion 2: Table 0 must be 1 row, 2 columns
        table = doc.tables[0]
        assert len(table.rows) == 1, f"[{doc_type}] Header table must have exactly 1 row, found {len(table.rows)}"
        assert len(table.columns) == 2, f"[{doc_type}] Header table must have exactly 2 columns, found {len(table.columns)}"

        # Assertion 3: Cell (0, 0) - Reference No, Left aligned
        cell_0 = table.cell(0, 0)
        c0_text = cell_0.text.strip()
        assert "Reference No:" in c0_text, f"[{doc_type}] Cell (0, 0) missing 'Reference No:', got: '{c0_text}'"
        
        # Check alignment of paragraph in cell 0
        p0 = cell_0.paragraphs[0]
        assert p0.alignment in (WD_ALIGN_PARAGRAPH.LEFT, None), \
            f"[{doc_type}] Cell (0, 0) paragraph should be LEFT aligned, got {p0.alignment}"

        # Assertion 4: Cell (0, 1) - Dated: or Date:, Right aligned
        cell_1 = table.cell(0, 1)
        c1_text = cell_1.text.strip()
        if doc_type in ["experience_certificate", "salary_certificate"]:
            assert "Date:" in c1_text, f"[{doc_type}] Cell (0, 1) expected 'Date:', got '{c1_text}'"
        else:
            assert "Dated:" in c1_text, f"[{doc_type}] Cell (0, 1) expected 'Dated:', got '{c1_text}'"

        p1 = cell_1.paragraphs[0]
        assert p1.alignment == WD_ALIGN_PARAGRAPH.RIGHT, \
            f"[{doc_type}] Cell (0, 1) paragraph should be RIGHT aligned, got {p1.alignment}"

        # Assertion 5: No old space-padding pattern in any paragraph
        for p in doc.paragraphs:
            assert (" " * 50) not in p.text, \
                f"[{doc_type}] Detected old space padding string in paragraph: '{p.text[:60]}...'"

        print(f"  ✅ '{doc_type}' passed: 2-column borderless header table with left/right alignment verified.")


if __name__ == "__main__":
    print("=" * 70)
    print("TEST WORD DATE ALIGNMENT (Task 3)")
    print("=" * 70)
    test_formatters()
    test_document_date_alignment()
    print("\n🎉 ALL WORD DATE ALIGNMENT TESTS PASSED SUCCESSFULLY!")
