"""
Verification Script: test_document_generation.py
Tests the Document Generation Orchestrator Service (Part 1):
1. Candidate creation / retrieval with full fields.
2. Preview generation for Offer letter (verifies reference number is NOT consumed).
3. Final document generation for Offer letter (verifies reference number consumed + audit log created).
4. Increment letter preview + final cycle with extra fields.
5. Termination letter preview + final cycle (generic template resolution).
6. Experience certificate with exit_date=None ('is currently working') and exit_date set ('was working').
7. Salary certificate generation.
8. Template resolution fallback when designation has no template (returns available_designations).
"""

import os
import sys
import datetime
import docx

# Ensure desktop_app is in sys.path
SCRIPT_DIR = os.path.dirname(os.path.abspath(__file__))
DESKTOP_APP_DIR = os.path.dirname(SCRIPT_DIR)
if DESKTOP_APP_DIR not in sys.path:
    sys.path.insert(0, DESKTOP_APP_DIR)

sys.stdout.reconfigure(encoding="utf-8")

from database import db_manager
from services import document_generation_service


def run_tests():
    print("=" * 70)
    print("PART 1 VERIFICATION: DOCUMENT GENERATION ORCHESTRATOR SERVICE")
    print("=" * 70)

    db_manager.init_db()

    # Ensure a valid reference configuration with sufficient room for test runs
    ref_cfg = db_manager.get_reference_config("hr_reference")
    if not ref_cfg or ref_cfg["current_number"] + 10 > ref_cfg["range_end"]:
        db_manager.save_reference_config("hr_reference", "DFY/HR/{YEAR}/{NUM}", 3, 1, 9999)
        ref_cfg = db_manager.get_reference_config("hr_reference")
    print(f"Initial Reference Config: {ref_cfg}")

    # Step 1: Create / update test candidate
    test_mobile = "9876543210"
    cand_data = {
        "candidate_name": "Rohan Sharma",
        "mobile": test_mobile,
        "email": "rohan.sharma@example.com",
        "applied_designation": "Data Entry Operator",
        "negotiated_salary": "28000",
        "recent_designation": "DEO",
        "highest_education": "Graduate",
        "total_experience": "3 Years",
        "core_technical_skills": "MS Excel, Typing",
        "key_responsibilities": "Data management",
        "contact_location": "New Delhi",
    }
    db_manager.upsert_candidate(cand_data, "test_joining_sheet.xlsx")
    cand = db_manager.get_candidate_by_mobile(test_mobile)
    assert cand is not None, "Candidate upsert failed"
    candidate_id = cand["id"]

    # Populate extra candidate details
    db_manager.update_candidate_extra_details(
        candidate_id=candidate_id,
        father_name="Suresh Sharma",
        full_address="Flat 102, Green Park Apartments, Sector 15, Rohini, Delhi",
        gender="Male",
        department="Operations",
        date_joined="2024-04-01",
        employee_number="DFY-7788"
    )
    cand = db_manager.get_candidate_by_id(candidate_id)
    print(f"1. Test candidate prepared: ID {candidate_id}, Name: {cand['candidate_name']}, Desig: {cand['applied_designation']}")
    assert cand["father_name"] == "Suresh Sharma"
    print("   ✅ Candidate details verified.")

    # Step 2: Offer letter PREVIEW (Reference number must NOT increment)
    curr_num_before_preview = db_manager.get_reference_config("hr_reference")["current_number"]
    preview_res = document_generation_service.generate_preview_document(
        candidate_id=candidate_id,
        document_type="offer"
    )
    curr_num_after_preview = db_manager.get_reference_config("hr_reference")["current_number"]

    print(f"\n2. Offer Preview Result: success={preview_res.get('success')}")
    print(f"   Preview file path: {preview_res.get('preview_file_path')}")
    print(f"   Previewed ref number: {preview_res.get('previewed_reference_number')}")
    assert preview_res["success"] is True, f"Offer preview failed: {preview_res}"
    assert os.path.exists(preview_res["preview_file_path"]), "Preview file does not exist on disk"
    assert curr_num_before_preview == curr_num_after_preview, \
        f"Preview consumed reference sequence! Before: {curr_num_before_preview}, After: {curr_num_after_preview}"
    print("   ✅ Offer preview verified (reference sequence was NOT consumed).")

    # Step 3: Offer letter FINALIZE (Reference number MUST increment & audit log created)
    curr_num_before_final = db_manager.get_reference_config("hr_reference")["current_number"]
    final_res = document_generation_service.finalize_document(
        candidate_id=candidate_id,
        document_type="offer"
    )
    curr_num_after_final = db_manager.get_reference_config("hr_reference")["current_number"]

    print(f"\n3. Offer Finalize Result: success={final_res.get('success')}")
    print(f"   Final file path: {final_res.get('file_path')}")
    print(f"   Assigned ref number: {final_res.get('reference_number')}")
    print(f"   Log ID: {final_res.get('log_id')}")
    assert final_res["success"] is True, f"Offer finalize failed: {final_res}"
    assert os.path.exists(final_res["file_path"]), "Final document file does not exist on disk"
    assert curr_num_after_final == curr_num_before_final + 1, \
        f"Reference sequence did not increment! Before: {curr_num_before_final}, After: {curr_num_after_final}"

    # Verify audit log
    history = db_manager.get_generation_history(candidate_id)
    assert len(history) > 0, "No audit log found in document_generation_log"
    assert history[0]["reference_number"] == final_res["reference_number"]
    assert history[0]["document_type"] == "offer"
    print("   ✅ Offer final generation verified (reference sequence incremented + audit log created).")

    # Step 4: Increment letter PREVIEW + FINALIZE with extra fields
    inc_extra = {
        "effective_from": "2025-04-01",
        "revised_pay_grade": "32000",
        "offer_letter_document_date": "2024-03-25"
    }
    inc_preview = document_generation_service.generate_preview_document(
        candidate_id=candidate_id,
        document_type="increment",
        extra_fields=inc_extra
    )
    assert inc_preview["success"] is True, f"Increment preview failed: {inc_preview}"

    inc_final = document_generation_service.finalize_document(
        candidate_id=candidate_id,
        document_type="increment",
        extra_fields=inc_extra
    )
    assert inc_final["success"] is True, f"Increment finalize failed: {inc_final}"
    assert os.path.exists(inc_final["file_path"]), "Increment final file missing"
    print(f"\n4. Increment letter finalized successfully: {inc_final['reference_number']}")
    print("   ✅ Increment letter cycle verified.")

    # Step 5: Termination letter PREVIEW + FINALIZE (uses Generic template automatically)
    term_extra = {
        "termination_date": "2025-05-31",
        "reason": "Project closure at location",
        "notice_period": "30 days"
    }
    term_preview = document_generation_service.generate_preview_document(
        candidate_id=candidate_id,
        document_type="termination",
        extra_fields=term_extra
    )
    assert term_preview["success"] is True, f"Termination preview failed: {term_preview}"

    term_final = document_generation_service.finalize_document(
        candidate_id=candidate_id,
        document_type="termination",
        extra_fields=term_extra
    )
    assert term_final["success"] is True, f"Termination finalize failed: {term_final}"
    assert os.path.exists(term_final["file_path"]), "Termination final file missing"
    print(f"\n5. Termination letter finalized successfully: {term_final['reference_number']}")
    print("   ✅ Termination letter cycle verified.")

    # Step 6: Experience certificate (exit_date=None -> 'is currently working' vs exit_date set -> 'was working')
    # 6a. Active employee (exit_date=None)
    exp_active = document_generation_service.finalize_document(
        candidate_id=candidate_id,
        document_type="experience",
        extra_fields={"exit_date": None}
    )
    assert exp_active["success"] is True
    doc_active = docx.Document(exp_active["file_path"])
    full_text_active = "\n".join(p.text for p in doc_active.paragraphs)
    assert "is currently working" in full_text_active, "Active certificate text missing 'is currently working'"
    print("\n6a. Experience certificate (Active / Till Date) verified: contains 'is currently working'")

    # 6b. Relieved employee (exit_date set)
    exp_relieved = document_generation_service.finalize_document(
        candidate_id=candidate_id,
        document_type="experience",
        extra_fields={"exit_date": "2025-06-30"}
    )
    assert exp_relieved["success"] is True
    doc_relieved = docx.Document(exp_relieved["file_path"])
    full_text_relieved = "\n".join(p.text for p in doc_relieved.paragraphs)
    assert "was working" in full_text_relieved, "Relieved certificate text missing 'was working'"
    print("6b. Experience certificate (Relieved) verified: contains 'was working'")
    print("   ✅ Both experience certificate phrasing variants verified.")

    # Step 7: Salary certificate test
    salary_cert = document_generation_service.finalize_document(
        candidate_id=candidate_id,
        document_type="salary",
        extra_fields={"exit_date": None}
    )
    assert salary_cert["success"] is True
    doc_salary = docx.Document(salary_cert["file_path"])
    full_text_salary = "\n".join(p.text for p in doc_salary.paragraphs)
    assert "drawing a total monthly payment of INR" in full_text_salary
    print(f"\n7. Salary certificate finalized successfully: {salary_cert['reference_number']}")
    print("   ✅ Salary certificate text verified.")

    # Step 8: Template resolution fallback (unmatched designation)
    dummy_cand = {
        "candidate_name": "Anoop Verma",
        "applied_designation": "Astronaut Special Coordinator",
    }
    res_unmatched = document_generation_service.resolve_template("offer", dummy_cand)
    print(f"\n8. Unmatched designation resolution result: found={res_unmatched['found']}")
    print(f"   Available designations count: {len(res_unmatched['available_designations'])}")
    assert res_unmatched["found"] is False, "Unmatched designation should yield found=False"
    assert len(res_unmatched["available_designations"]) > 0, "available_designations should not be empty"
    print("   ✅ Fallback available_designations list verified.")

    print("\n" + "=" * 70)
    print("ALL PART 1 ORCHESTRATOR TESTS PASSED SUCCESSFULLY! ✅")
    print("=" * 70)


if __name__ == "__main__":
    run_tests()
