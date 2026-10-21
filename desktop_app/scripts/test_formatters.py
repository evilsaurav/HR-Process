"""
Verification Script: test_formatters.py
Tests:
1. All document_formatters utility functions with multiple test cases.
2. Robust placeholder replacement on real Offer letter template (double-braces + split runs + multiline address).
3. Robust placeholder replacement on real Increment letter template (single-braces {Job Title} etc.).
"""

import os
import sys
import datetime
import docx

# Configure UTF-8 encoding for console output on Windows
sys.stdout.reconfigure(encoding="utf-8")

# Ensure desktop_app is in sys.path
SCRIPT_DIR = os.path.dirname(os.path.abspath(__file__))
DESKTOP_APP_DIR = os.path.dirname(SCRIPT_DIR)
if DESKTOP_APP_DIR not in sys.path:
    sys.path.insert(0, DESKTOP_APP_DIR)

from services.document_formatters import (
    get_title_and_relation,
    format_address_multiline,
    amount_to_indian_words,
    format_date_long,
    calculate_end_date,
    format_currency_inr
)
from services.placeholder_engine import replace_placeholders_in_docx
from config import TEMPLATES_DIR


def test_formatters():
    print("=" * 70)
    print("TEST SUITE 1: DOCUMENT FORMATTERS UTILITIES")
    print("=" * 70)

    # 1. get_title_and_relation
    print("\n[1] get_title_and_relation:")
    cases_gender = [
        ("Male", ("Mr.", "S/o")),
        ("m", ("Mr.", "S/o")),
        ("Female", ("Ms.", "D/o")),
        ("f", ("Ms.", "D/o")),
        ("", ("Mr.", "S/o")),
        (None, ("Mr.", "S/o")),
        ("Unknown", ("Mr.", "S/o"))
    ]
    for inp, expected in cases_gender:
        actual = get_title_and_relation(inp)
        status = "PASS" if actual == expected else "FAIL"
        print(f"  Input: {repr(inp):15} -> Expected: {expected} | Actual: {actual} [{status}]")
        assert actual == expected

    # 2. format_address_multiline
    print("\n[2] format_address_multiline:")
    # Short address (<= 3 parts)
    addr_short = "Flat 101, Galaxy Tower, Patna"
    expected_short = "Flat 101,\nGalaxy Tower,\nPatna"
    actual_short = format_address_multiline(addr_short)
    print(f"  Short address lines: {len(actual_short.splitlines())} lines:\n    {repr(actual_short)}")
    assert len(actual_short.splitlines()) == 3

    # Long address (> 3 parts)
    addr_long = "Flat 101, Galaxy Heights, Near Metro Station, Sector 62, Noida, UP, 201301"
    actual_long = format_address_multiline(addr_long)
    print(f"  Long address packed into max 3 lines:\n    Line 1: {repr(actual_long.splitlines()[0])}\n    Line 2: {repr(actual_long.splitlines()[1])}\n    Line 3: {repr(actual_long.splitlines()[2])}")
    assert len(actual_long.splitlines()) == 3

    # 3. amount_to_indian_words
    print("\n[3] amount_to_indian_words:")
    cases_amount = [
        (25000, "Twenty-Five Thousand Only"),
        ("25,000", "Twenty-Five Thousand Only"),
        ("₹ 1,25,000/-", "One Lakh, Twenty-Five Thousand Only"),
        (0, "Zero Only")
    ]
    for inp, expected in cases_amount:
        actual = amount_to_indian_words(inp)
        status = "PASS" if actual == expected else "FAIL"
        print(f"  Input: {repr(inp):18} -> Actual: {repr(actual)} [{status}]")
        assert actual == expected

    # 4. format_date_long
    print("\n[4] format_date_long:")
    cases_dates = [
        ("2024-01-15", "15 January 2024"),
        ("15/01/2024", "15 January 2024"),
        ("15-01-2024", "15 January 2024"),
        (datetime.date(2024, 1, 15), "15 January 2024"),
        ("InvalidDate", "InvalidDate"),
        (None, "")
    ]
    for inp, expected in cases_dates:
        actual = format_date_long(inp)
        status = "PASS" if actual == expected else "FAIL"
        print(f"  Input: {repr(inp):22} -> Actual: {repr(actual)} [{status}]")
        assert actual == expected

    # 5. calculate_end_date
    print("\n[5] calculate_end_date:")
    cases_end = [
        ("15 January 2024", "14 January 2025"),
        ("2024-01-15", "14 January 2025"),
        ("2024-02-29", "27 February 2025")  # Leap year rollover
    ]
    for inp, expected in cases_end:
        actual = calculate_end_date(inp)
        status = "PASS" if actual == expected else "FAIL"
        print(f"  Start: {repr(inp):18} -> 1 Yr End: {repr(actual)} [{status}]")
        assert actual == expected

    # 6. format_currency_inr
    print("\n[6] format_currency_inr:")
    cases_curr = [
        (25000, "₹25,000"),
        ("35000", "₹35,000"),
        ("₹ 15,500/-", "₹15,500")
    ]
    for inp, expected in cases_curr:
        actual = format_currency_inr(inp)
        status = "PASS" if actual == expected else "FAIL"
        print(f"  Input: {repr(inp):16} -> Actual: {repr(actual)} [{status}]")
        assert actual == expected


def test_offer_template_replacement():
    print("\n" + "=" * 70)
    print("TEST SUITE 2: PLACEHOLDER ENGINE ON REAL OFFER TEMPLATE")
    print("=" * 70)

    template_file = os.path.join(TEMPLATES_DIR, "offer", "associate_district_coordinator.docx")
    print(f"Loading template: {template_file}")
    doc = docx.Document(template_file)

    multiline_address = format_address_multiline("House No 42, Green Park Avenue, Near Gandhi Maidan, Patna, Bihar, 800001")

    mapping = {
        "TITLE": "Mr.",
        "Full Name": "Rahul Sharma",
        "RELATION": "S/o",
        "FATHER_TITLE": "Mr.",
        "Father Name": "Sunil Sharma",
        "Full Address": multiline_address,
        "Job Title": "Associate District Coordinator",
        "Department": "PPSA Bihar",
        "Location": "Patna",
        "Date Joined": "15 January 2024",
        "END DATE": calculate_end_date("15 January 2024"),
        "Pay Grade": format_currency_inr(35000),
        "SALARY_WORDS": amount_to_indian_words(35000)
    }

    print(f"Replacing {len(mapping)} placeholder keys across all runs and paragraphs...")
    replace_placeholders_in_docx(doc, mapping)

    output_dir = os.path.join(DESKTOP_APP_DIR, "output")
    os.makedirs(output_dir, exist_ok=True)
    output_path = os.path.join(output_dir, "test_offer_output.docx")
    doc.save(output_path)
    print(f"Saved generated test offer document to: {output_path}")

    # Re-read and inspect paragraphs
    doc_verify = docx.Document(output_path)
    print("\nRe-read paragraphs from generated file:")
    unresolved_braces = []
    for i, p in enumerate(doc_verify.paragraphs):
        text = p.text.strip()
        if text:
            if "{{" in text or "}}" in text:
                unresolved_braces.append((i, text))
            if any(k in text for k in ["Rahul Sharma", "Associate District Coordinator", "PPSA Bihar", "Thirty-Five Thousand"]):
                print(f"  [P{i:02d} MATCH]: {text}")

    # Check multi-line address paragraph
    p_addr = doc_verify.paragraphs[7]
    print(f"\nAddress paragraph (P7): {repr(p_addr.text)}")
    assert "Rahul Sharma" in doc_verify.paragraphs[5].text
    assert "Sunil Sharma" in doc_verify.paragraphs[6].text
    assert "\n" in p_addr.text
    print(f"Unresolved placeholders found in offer letter: {len(unresolved_braces)}")
    assert len(unresolved_braces) == 0, f"Unresolved placeholders: {unresolved_braces}"
    print("Offer letter placeholder replacement: 100% SUCCESS")


def test_increment_template_replacement():
    print("\n" + "=" * 70)
    print("TEST SUITE 3: SINGLE-BRACE REPLACEMENT ON INCREMENT TEMPLATE")
    print("=" * 70)

    template_file = os.path.join(TEMPLATES_DIR, "increment", "block_coordinator.docx")
    print(f"Loading template: {template_file}")
    doc = docx.Document(template_file)

    mapping = {
        "TITLE": "Mr.",
        "Full Name": "Vikas Gupta",
        "RELATION": "S/o",
        "FATHER_TITLE": "Mr.",
        "Father Name": "Ramesh Gupta",
        "Full Address": format_address_multiline("Lane 5, Boring Road, Near City Mall, Patna, Bihar, 800001"),
        "Job Title": "Block Coordinator",
        "Effective From": "01 April 2024",
        "Pay Grade": "20,000",
        "Revised Pay Grade": "25,000",
        "Offer letter Document Date": "01 January 2023"
    }

    replace_placeholders_in_docx(doc, mapping)

    output_dir = os.path.join(DESKTOP_APP_DIR, "output")
    output_path = os.path.join(output_dir, "test_increment_output.docx")
    doc.save(output_path)
    print(f"Saved generated test increment document to: {output_path}")

    # Re-read and inspect
    doc_verify = docx.Document(output_path)
    unresolved = []
    for i, p in enumerate(doc_verify.paragraphs):
        text = p.text.strip()
        if text:
            # Check for any remaining '{' or '}' placeholders
            if "{" in text or "}" in text:
                unresolved.append((i, text))
            if any(k in text for k in ["Vikas Gupta", "Block Coordinator", "01 April 2024", "25,000"]):
                print(f"  [P{i:02d} MATCH]: {text}")

    print(f"Unresolved placeholders found in increment letter: {len(unresolved)}")
    assert len(unresolved) == 0, f"Unresolved single-brace placeholders: {unresolved}"
    print("Increment letter single-brace placeholder replacement: 100% SUCCESS")


if __name__ == "__main__":
    test_formatters()
    test_offer_template_replacement()
    test_increment_template_replacement()
    print("\n" + "=" * 70)
    print("ALL TEST SUITES PASSED CLEANLY!")
    print("=" * 70)
